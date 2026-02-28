# Copyright (c) 2025 Bytedance Ltd. and/or its affiliates
# SPDX-License-Identifier: MIT

import json
import logging
import os
import re
import uuid
from pathlib import Path
from typing import Literal, Optional

from langchain_core.messages import HumanMessage, SystemMessage
from langchain_core.tools import StructuredTool
from pydantic import BaseModel, Field

from src.llms.llm import get_llm_by_type

logger = logging.getLogger(__name__)

PPT_OUTPUT_DIR = "ppt_output"
OUTLINE_SYSTEM = """你是一个专业的PPT大纲撰写助手。根据用户给出的一句话主题，生成一份结构清晰的PPT大纲。
要求：
1. 使用 Markdown 格式输出。
2. 第一行用一级标题（# ）写整份PPT的标题。
3. 每一页幻灯片用一个二级标题（## ）作为该页标题，紧跟的列表项（- 或 *）为该页要点，每页3-6个要点为宜。
4. 总页数控制在5-10页（含封面）。
5. 只输出大纲内容，不要输出其他解释。"""


class GeneratePPTInput(BaseModel):
    """Input for the PPT generation tool (outline or generate)."""

    topic: str = Field(default="", description="One-sentence topic for the PPT; required when action is 'outline'.")
    outline: str = Field(default="", description="Outline content; required when action is 'generate' (from step 1 or edited by user).")
    action: Literal["outline", "generate"] = Field(
        default="outline",
        description="'outline' to generate outline from topic; 'generate' to create .pptx from outline.",
    )


def _generate_outline(topic: str) -> str:
    """Use LLM to generate PPT outline from one-sentence topic."""
    if not (topic and topic.strip()):
        return json.dumps({"error": "topic is required when action is 'outline'"}, ensure_ascii=False)
    try:
        llm = get_llm_by_type("basic")
        messages = [
            SystemMessage(content=OUTLINE_SYSTEM),
            HumanMessage(content=f"请根据以下主题生成PPT大纲：\n\n{topic.strip()}"),
        ]
        response = llm.invoke(messages)
        outline_text = response.content if hasattr(response, "content") else str(response)
        outline_text = (outline_text or "").strip()
        return json.dumps({"outline": outline_text}, ensure_ascii=False)
    except Exception as e:
        logger.exception("PPT outline generation failed")
        return json.dumps({"error": f"生成大纲失败: {e!s}"}, ensure_ascii=False)


def _parse_outline_to_slides(outline: str) -> list[dict]:
    """Parse markdown outline into list of slides: [{"title": str, "bullets": [str]}, ...]."""
    lines = [s.strip() for s in outline.strip().split("\n") if s.strip()]
    slides = []
    current_title = None
    current_bullets = []

    for line in lines:
        if line.startswith("# "):
            title = line.lstrip("# ").strip()
            if not title:
                continue
            if current_title is not None:
                slides.append({"title": current_title, "bullets": current_bullets})
            current_title = title
            current_bullets = []
        elif line.startswith("## "):
            title = line.lstrip("# ").strip()
            if not title:
                continue
            if current_title is not None:
                slides.append({"title": current_title, "bullets": current_bullets})
            current_title = title
            current_bullets = []
        elif line.startswith("- ") or line.startswith("* "):
            bullet = line[2:].strip()
            if bullet:
                current_bullets.append(bullet)
        else:
            if current_title and line and not line.startswith("#"):
                current_bullets.append(line)

    if current_title is not None:
        slides.append({"title": current_title, "bullets": current_bullets})

    if not slides:
        first_line = lines[0] if lines else ""
        first_line = re.sub(r"^#+\s*", "", first_line).strip()
        if first_line:
            slides = [{"title": first_line, "bullets": []}]
            for line in lines[1:]:
                line = line.strip()
                if line and (line.startswith("- ") or line.startswith("* ")):
                    slides[0]["bullets"].append(line[2:].strip())
                elif line and not line.startswith("#"):
                    slides.append({"title": line, "bullets": []})
    return slides


def _generate_pptx_from_outline(outline: str) -> str:
    """Create .pptx from outline and return JSON with download_url and filename."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return json.dumps({"error": "python-pptx 未安装，请运行: uv add python-pptx"}, ensure_ascii=False)

    if not (outline and outline.strip()):
        return json.dumps({"error": "outline is required when action is 'generate'"}, ensure_ascii=False)

    slides_data = _parse_outline_to_slides(outline)
    if not slides_data:
        return json.dumps({"error": "无法从大纲解析出幻灯片内容，请检查格式（# 标题，## 页标题，- 要点）"}, ensure_ascii=False)

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    title_slide_layout = prs.slide_layouts[0]
    content_slide_layout = prs.slide_layouts[1]

    for i, slide_data in enumerate(slides_data):
        title = (slide_data.get("title") or "").strip() or f"Slide {i + 1}"
        bullets = slide_data.get("bullets") or []

        if i == 0 and len(slides_data) > 1 and not bullets:
            slide = prs.slides.add_slide(title_slide_layout)
            slide.shapes.title.text = title
            if slide.placeholders[1]:
                slide.placeholders[1].text = ""
        else:
            slide = prs.slides.add_slide(content_slide_layout)
            slide.shapes.title.text = title
            body = slide.placeholders[1].text_frame
            body.clear()
            for b in bullets[:8]:
                p = body.add_paragraph()
                p.text = b
                p.level = 0
                p.font.size = Pt(14)

    cwd = Path(os.getcwd()).resolve()
    out_dir = cwd / PPT_OUTPUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = f"ppt_{uuid.uuid4().hex[:12]}.pptx"
    filepath = out_dir / filename
    prs.save(str(filepath))
    relative_path = f"{PPT_OUTPUT_DIR}/{filename}"
    download_url = f"/api/workspace-file?path={relative_path}"
    return json.dumps(
        {"download_url": download_url, "filename": filename, "path": relative_path},
        ensure_ascii=False,
    )


def _run_generate_ppt(inp: GeneratePPTInput | dict) -> str:
    if isinstance(inp, dict):
        topic = inp.get("topic", "") or ""
        outline = inp.get("outline", "") or ""
        action = inp.get("action", "outline") or "outline"
    else:
        topic = getattr(inp, "topic", "") or ""
        outline = getattr(inp, "outline", "") or ""
        action = getattr(inp, "action", "outline") or "outline"
    if action == "outline":
        return _generate_outline(topic)
    if action == "generate":
        return _generate_pptx_from_outline(outline)
    return json.dumps({"error": f"invalid action: {action}"}, ensure_ascii=False)


generate_ppt_tool = StructuredTool(
    name="generate_ppt_tool",
    description="Generate a PPT: step 1 use action='outline' with topic to get outline; step 2 use action='generate' with outline to create .pptx and get download link.",
    args_schema=GeneratePPTInput,
    func=_run_generate_ppt,
)
